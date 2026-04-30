from __future__ import annotations

import argparse
import contextlib
import gc
import html
import json
import re
import shutil
import sys
import unicodedata
import zipfile
from collections import Counter, defaultdict
from dataclasses import asdict, dataclass
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable

import cv2
import numpy as np
import pandas as pd
from PIL import Image, ImageOps
from PIL.ExifTags import GPSTAGS, TAGS


IMAGE_EXTENSIONS = {".jpg", ".jpeg", ".png", ".bmp", ".tif", ".tiff", ".webp"}


@dataclass(frozen=True)
class TargetDefinition:
    label: str
    prompt: str
    group: str
    priority: str
    color: str
    action: str
    synonyms: tuple[str, ...] = ()


@dataclass(frozen=True)
class SearchPass:
    name: str
    kind: str
    max_side: int = 1600
    tile_size: int = 960
    overlap: int = 160


@dataclass(frozen=True)
class SearchCrop:
    image: np.ndarray
    x0: int
    y0: int
    scale: float
    pass_name: str


@dataclass
class Detection:
    image: str
    target: str
    prompt: str
    group: str
    priority: str
    confidence: float
    x1: int
    y1: int
    x2: int
    y2: int
    width_px: int
    height_px: int
    area_ratio: float
    source_pass: str
    gps_lat: float | None = None
    gps_lon: float | None = None
    gps_alt_m: float | None = None
    taken_at: str | None = None


@dataclass
class ImageSummary:
    image: str
    status: str
    total_detections: int
    findings: str
    original_width: int
    original_height: int
    processed_width: int
    processed_height: int
    annotated_path: str | None = None
    gps_lat: float | None = None
    gps_lon: float | None = None
    gps_alt_m: float | None = None
    taken_at: str | None = None


COMMON_ACTIONS = {
    "EPP": "Verificar dotacion, uso y registro de EPP antes de liberar el frente.",
    "Senalizacion": "Revisar delimitacion, carteleria visible y continuidad del vallado.",
    "Vehiculos": "Separar circulacion peatonal/equipos y confirmar condiciones de maniobra.",
    "Maquinaria": "Validar estado operativo, radios de giro y distancia segura con personal.",
    "Ductos": "Contrastar contra traza, interferencias, canerias, mojones y puntos de control.",
    "Excavaciones": "Revisar accesos, estabilidad, vallado, distancia a bordes y servicios existentes.",
    "Altura": "Confirmar plataforma, linea de vida, barandas y planilla de control de equipos.",
    "Orden y limpieza": "Registrar condicion, retirar obstaculos y normalizar acopios o residuos.",
}


CATALOG = (
    TargetDefinition(
        "Cono vial",
        "traffic cone",
        "Senalizacion",
        "Media",
        "#f97316",
        COMMON_ACTIONS["Senalizacion"],
        ("cono", "conos", "cono vial", "traffic cone", "baliza"),
    ),
    TargetDefinition(
        "Cartel de seguridad",
        "safety sign",
        "Senalizacion",
        "Alta",
        "#eab308",
        COMMON_ACTIONS["Senalizacion"],
        ("cartel", "carteles", "senal", "senalizacion", "safety sign", "warning sign"),
    ),
    TargetDefinition(
        "Casco",
        "hard hat",
        "EPP",
        "Alta",
        "#22c55e",
        COMMON_ACTIONS["EPP"],
        ("casco", "cascos", "hardhat", "hard hat", "helmet", "casco de seguridad"),
    ),
    TargetDefinition(
        "Chaleco reflectivo",
        "safety vest",
        "EPP",
        "Alta",
        "#84cc16",
        COMMON_ACTIONS["EPP"],
        ("chaleco", "chalecos", "chaleco reflectivo", "safety vest", "reflective vest"),
    ),
    TargetDefinition(
        "Persona",
        "person",
        "EPP",
        "Media",
        "#14b8a6",
        "Cruzar presencia de personal con EPP visible y segregacion de equipos.",
        ("persona", "personas", "operario", "trabajador", "worker", "person"),
    ),
    TargetDefinition(
        "Auto",
        "car",
        "Vehiculos",
        "Media",
        "#2563eb",
        COMMON_ACTIONS["Vehiculos"],
        ("auto", "autos", "coche", "car", "cars", "vehiculo liviano"),
    ),
    TargetDefinition(
        "Camioneta",
        "pickup truck",
        "Vehiculos",
        "Media",
        "#1d4ed8",
        COMMON_ACTIONS["Vehiculos"],
        ("camioneta", "pickup", "pickup truck", "4x4", "utilitario"),
    ),
    TargetDefinition(
        "Camion",
        "truck",
        "Vehiculos",
        "Media",
        "#0f766e",
        COMMON_ACTIONS["Vehiculos"],
        ("camion", "camiones", "truck", "dump truck", "volcador"),
    ),
    TargetDefinition(
        "Excavadora",
        "excavator",
        "Maquinaria",
        "Alta",
        "#a855f7",
        COMMON_ACTIONS["Maquinaria"],
        ("excavadora", "retroexcavadora", "excavator", "backhoe"),
    ),
    TargetDefinition(
        "Grua",
        "crane",
        "Maquinaria",
        "Alta",
        "#7c3aed",
        COMMON_ACTIONS["Maquinaria"],
        ("grua", "crane", "hidrogrua", "sideboom"),
    ),
    TargetDefinition(
        "Caneria/Tuberia",
        "pipe",
        "Ductos",
        "Media",
        "#0891b2",
        COMMON_ACTIONS["Ductos"],
        ("caneria", "canerias", "tuberia", "tuberias", "pipe", "pipeline", "ducto"),
    ),
    TargetDefinition(
        "Zanja abierta",
        "trench",
        "Excavaciones",
        "Alta",
        "#b45309",
        COMMON_ACTIONS["Excavaciones"],
        ("zanja", "zanjas", "excavacion", "excavaciones", "trench", "ditch"),
    ),
    TargetDefinition(
        "Mojon/estaca",
        "survey stake",
        "Ductos",
        "Media",
        "#dc2626",
        COMMON_ACTIONS["Ductos"],
        ("mojon", "mojones", "estaca", "estacas", "survey stake", "survey marker"),
    ),
    TargetDefinition(
        "Escalera",
        "ladder",
        "Altura",
        "Alta",
        "#f43f5e",
        COMMON_ACTIONS["Altura"],
        ("escalera", "escaleras", "ladder"),
    ),
    TargetDefinition(
        "Andamio",
        "scaffolding",
        "Altura",
        "Alta",
        "#be123c",
        COMMON_ACTIONS["Altura"],
        ("andamio", "andamios", "scaffold", "scaffolding"),
    ),
    TargetDefinition(
        "Residuo u obstaculo",
        "construction debris",
        "Orden y limpieza",
        "Media",
        "#64748b",
        COMMON_ACTIONS["Orden y limpieza"],
        ("residuo", "residuos", "escombro", "escombros", "obstaculo", "debris"),
    ),
)


DEFAULT_PRESET_LABELS = (
    "Cono vial",
    "Cartel de seguridad",
    "Casco",
    "Chaleco reflectivo",
    "Persona",
    "Auto",
    "Camioneta",
    "Camion",
    "Excavadora",
    "Caneria/Tuberia",
    "Zanja abierta",
    "Mojon/estaca",
)


SEARCH_PROFILES = {
    "rapida": (
        SearchPass("general_1280", "global", max_side=1280),
    ),
    "equilibrada": (
        SearchPass("general_1600", "global", max_side=1600),
        SearchPass("tiles_960", "tile", tile_size=960, overlap=180),
    ),
    "profunda": (
        SearchPass("general_1800", "global", max_side=1800),
        SearchPass("tiles_960", "tile", tile_size=960, overlap=220),
        SearchPass("tiles_640", "tile", tile_size=640, overlap=180),
    ),
}


def normalize_key(value: str) -> str:
    value = unicodedata.normalize("NFKD", str(value))
    value = value.encode("ascii", "ignore").decode("ascii")
    value = re.sub(r"[^a-zA-Z0-9]+", " ", value.lower()).strip()
    return re.sub(r"\s+", " ", value)


def split_terms(value: str | Iterable[str] | None) -> list[str]:
    if value is None:
        return []
    if isinstance(value, str):
        parts = re.split(r"[,;\n]+", value)
    else:
        parts = list(value)
    return [str(part).strip() for part in parts if str(part).strip()]


def catalog_lookup() -> dict[str, TargetDefinition]:
    lookup: dict[str, TargetDefinition] = {}
    for target in CATALOG:
        keys = (target.label, target.prompt, *target.synonyms)
        for key in keys:
            lookup[normalize_key(key)] = target
    return lookup


def build_targets(
    preset_labels: Iterable[str] | None = None,
    extra_terms: str | Iterable[str] | None = None,
) -> list[TargetDefinition]:
    lookup = catalog_lookup()
    selected: list[TargetDefinition] = []

    for label in preset_labels or ():
        target = lookup.get(normalize_key(label))
        if target is not None:
            selected.append(target)

    for term in split_terms(extra_terms):
        target = lookup.get(normalize_key(term))
        if target is None:
            clean = term.strip()
            target = TargetDefinition(
                label=clean.title(),
                prompt=clean,
                group="Busqueda personalizada",
                priority="Media",
                color="#475569",
                action="Validar visualmente el hallazgo y clasificarlo en el parte tecnico.",
                synonyms=(clean,),
            )
        selected.append(target)

    deduped: dict[str, TargetDefinition] = {}
    for target in selected:
        deduped[normalize_key(target.prompt)] = target
    return list(deduped.values())


def image_files(input_folder: str | Path) -> list[Path]:
    folder = Path(input_folder)
    files = [path for path in folder.iterdir() if path.suffix.lower() in IMAGE_EXTENSIONS]
    return sorted(files, key=lambda path: path.name.lower())


def hex_to_bgr(hex_color: str) -> tuple[int, int, int]:
    cleaned = hex_color.strip().lstrip("#")
    if len(cleaned) != 6:
        return (40, 120, 220)
    r = int(cleaned[0:2], 16)
    g = int(cleaned[2:4], 16)
    b = int(cleaned[4:6], 16)
    return (b, g, r)


def rational_to_float(value: Any) -> float | None:
    try:
        if isinstance(value, tuple) and len(value) == 2:
            return float(value[0]) / float(value[1])
        return float(value)
    except Exception:
        return None


def gps_coord_to_decimal(values: Any, ref: str | None) -> float | None:
    if not values or len(values) < 3:
        return None
    degrees = rational_to_float(values[0])
    minutes = rational_to_float(values[1])
    seconds = rational_to_float(values[2])
    if degrees is None or minutes is None or seconds is None:
        return None
    result = degrees + minutes / 60 + seconds / 3600
    if ref in {"S", "W"}:
        result *= -1
    return result


def extract_exif_metadata(path: str | Path) -> dict[str, Any]:
    metadata: dict[str, Any] = {"gps_lat": None, "gps_lon": None, "gps_alt_m": None, "taken_at": None}
    try:
        with Image.open(path) as image:
            exif = image.getexif()
            if not exif:
                return metadata
            decoded = {TAGS.get(tag, tag): value for tag, value in exif.items()}
            taken_at = decoded.get("DateTimeOriginal") or decoded.get("DateTime")
            if taken_at:
                metadata["taken_at"] = str(taken_at)

            gps_info = decoded.get("GPSInfo")
            if gps_info:
                gps = {GPSTAGS.get(key, key): value for key, value in gps_info.items()}
                metadata["gps_lat"] = gps_coord_to_decimal(gps.get("GPSLatitude"), gps.get("GPSLatitudeRef"))
                metadata["gps_lon"] = gps_coord_to_decimal(gps.get("GPSLongitude"), gps.get("GPSLongitudeRef"))
                altitude = rational_to_float(gps.get("GPSAltitude"))
                if altitude is not None and gps.get("GPSAltitudeRef") == 1:
                    altitude *= -1
                metadata["gps_alt_m"] = altitude
    except Exception:
        return metadata
    return metadata


def read_image_bgr(path: str | Path) -> tuple[np.ndarray, int, int]:
    with Image.open(path) as pil_image:
        pil_image = ImageOps.exif_transpose(pil_image).convert("RGB")
        width, height = pil_image.size
        rgb = np.array(pil_image)
    return cv2.cvtColor(rgb, cv2.COLOR_RGB2BGR), width, height


def resize_to_max_side(image: np.ndarray, max_side: int) -> tuple[np.ndarray, float]:
    height, width = image.shape[:2]
    longest = max(height, width)
    if longest <= max_side:
        return image, 1.0
    scale = max_side / float(longest)
    resized = cv2.resize(image, (int(width * scale), int(height * scale)), interpolation=cv2.INTER_AREA)
    return resized, scale


def starts_for_tiles(length: int, tile_size: int, overlap: int) -> list[int]:
    if length <= tile_size:
        return [0]
    step = max(1, tile_size - overlap)
    starts = list(range(0, length - tile_size + 1, step))
    last = max(0, length - tile_size)
    if not starts or starts[-1] != last:
        starts.append(last)
    return sorted(set(starts))


def build_search_crops(image: np.ndarray, passes: Iterable[SearchPass]) -> list[SearchCrop]:
    crops: list[SearchCrop] = []
    height, width = image.shape[:2]
    for search_pass in passes:
        if search_pass.kind == "global":
            resized, scale = resize_to_max_side(image, search_pass.max_side)
            crops.append(SearchCrop(resized, 0, 0, scale, search_pass.name))
            continue

        if search_pass.kind != "tile":
            continue

        tile = min(search_pass.tile_size, max(height, width))
        if height <= tile and width <= tile and any(crop.x0 == 0 and crop.y0 == 0 and crop.scale == 1.0 for crop in crops):
            continue
        x_starts = starts_for_tiles(width, tile, search_pass.overlap)
        y_starts = starts_for_tiles(height, tile, search_pass.overlap)
        for y0 in y_starts:
            for x0 in x_starts:
                crop = image[y0 : min(y0 + tile, height), x0 : min(x0 + tile, width)]
                crops.append(SearchCrop(crop, x0, y0, 1.0, search_pass.name))
    return crops


def safe_float(value: Any) -> float:
    try:
        if hasattr(value, "detach"):
            value = value.detach().cpu().numpy()
        if isinstance(value, np.ndarray):
            return float(value.reshape(-1)[0])
        return float(value)
    except Exception:
        return 0.0


def safe_xyxy(value: Any) -> tuple[float, float, float, float]:
    if hasattr(value, "detach"):
        value = value.detach().cpu().numpy()
    array = np.array(value).reshape(-1)
    return tuple(float(v) for v in array[:4])  # type: ignore[return-value]


def clamp_box(x1: float, y1: float, x2: float, y2: float, width: int, height: int) -> tuple[int, int, int, int]:
    ix1 = max(0, min(width - 1, int(round(x1))))
    iy1 = max(0, min(height - 1, int(round(y1))))
    ix2 = max(0, min(width - 1, int(round(x2))))
    iy2 = max(0, min(height - 1, int(round(y2))))
    if ix2 <= ix1:
        ix2 = min(width - 1, ix1 + 1)
    if iy2 <= iy1:
        iy2 = min(height - 1, iy1 + 1)
    return ix1, iy1, ix2, iy2


class VisualInspectionTool:
    def __init__(
        self,
        input_folder: str | Path,
        output_folder: str | Path = "output_reporte_visual",
        targets: Iterable[TargetDefinition] | None = None,
        confidence: float = 0.08,
        search_profile: str = "equilibrada",
        nms_iou: float = 0.45,
        max_image_side: int = 5200,
        batch_size: int = 8,
        imgsz: int = 960,
        model_name: str = "yolov8s-world.pt",
    ) -> None:
        self.input_folder = Path(input_folder)
        self.output_folder = Path(output_folder)
        self.annotated_folder = self.output_folder / "imagenes_marcadas"
        self.tables_folder = self.output_folder / "tablas"
        self.output_folder.mkdir(parents=True, exist_ok=True)
        self.annotated_folder.mkdir(parents=True, exist_ok=True)
        self.tables_folder.mkdir(parents=True, exist_ok=True)

        self.targets = list(targets or build_targets(DEFAULT_PRESET_LABELS))
        if not self.targets:
            raise ValueError("No hay objetivos de busqueda configurados.")

        self.prompts = [target.prompt for target in self.targets]
        self.prompt_to_target = {normalize_key(target.prompt): target for target in self.targets}
        self.confidence = float(confidence)
        self.search_profile = search_profile if search_profile in SEARCH_PROFILES else "equilibrada"
        self.search_passes = SEARCH_PROFILES[self.search_profile]
        self.nms_iou = float(nms_iou)
        self.max_image_side = int(max_image_side)
        self.batch_size = max(1, int(batch_size))
        self.imgsz = int(imgsz)
        self.model_name = model_name
        self.model = None
        self.detections: list[Detection] = []
        self.image_summaries: list[ImageSummary] = []
        self.started_at = datetime.now()
        self.finished_at: datetime | None = None

    def load_model(self) -> None:
        if self.model is not None:
            return
        try:
            from ultralytics import YOLOWorld
        except ImportError as exc:
            raise RuntimeError(
                "Falta ultralytics. En Colab ejecuta la celda de instalacion; localmente usa: "
                "pip install ultralytics opencv-python pillow reportlab pandas openpyxl"
            ) from exc

        self.model = YOLOWorld(self.model_name)
        self.model.set_classes(self.prompts)

    def process(self) -> list[ImageSummary]:
        self.load_model()
        files = image_files(self.input_folder)
        if not files:
            raise FileNotFoundError(f"No se encontraron imagenes en {self.input_folder}")

        print(f"Imagenes a procesar: {len(files)}")
        print(f"Objetivos simultaneos: {', '.join(target.label for target in self.targets)}")
        print(f"Perfil de busqueda: {self.search_profile} ({len(self.search_passes)} pasada/s)")

        for index, image_path in enumerate(files, start=1):
            print(f"[{index}/{len(files)}] {image_path.name}")
            summary = self.process_image(image_path)
            self.image_summaries.append(summary)
            gc.collect()

        self.finished_at = datetime.now()
        return self.image_summaries

    def process_image(self, image_path: Path) -> ImageSummary:
        metadata = extract_exif_metadata(image_path)
        image, original_width, original_height = read_image_bgr(image_path)
        image, process_scale = resize_to_max_side(image, self.max_image_side)
        processed_height, processed_width = image.shape[:2]
        crops = build_search_crops(image, self.search_passes)
        raw_detections = self.predict_crops(crops, image_path.name, processed_width, processed_height, metadata)
        detections = self.apply_nms(raw_detections)

        if process_scale != 1.0:
            note = f"procesada a {processed_width}x{processed_height}px para optimizar memoria"
        else:
            note = f"procesada a resolucion original {processed_width}x{processed_height}px"
        print(f"  sectores: {len(crops)} | detecciones: {len(detections)} | {note}")

        self.detections.extend(detections)
        counts = Counter(det.target for det in detections)
        findings = "; ".join(f"{label}: {count}" for label, count in counts.most_common()) or "Sin hallazgos"
        annotated_path: str | None = None

        if detections:
            annotated = self.draw_detections(image.copy(), image_path.name, detections)
            annotated_path = str(self.annotated_folder / f"{image_path.stem}_marcada.jpg")
            cv2.imwrite(annotated_path, annotated, [int(cv2.IMWRITE_JPEG_QUALITY), 92])

        return ImageSummary(
            image=image_path.name,
            status="Con hallazgos" if detections else "Sin hallazgos",
            total_detections=len(detections),
            findings=findings,
            original_width=original_width,
            original_height=original_height,
            processed_width=processed_width,
            processed_height=processed_height,
            annotated_path=annotated_path,
            **metadata,
        )

    def predict_crops(
        self,
        crops: list[SearchCrop],
        image_name: str,
        image_width: int,
        image_height: int,
        metadata: dict[str, Any],
    ) -> list[Detection]:
        detections: list[Detection] = []
        if self.model is None:
            raise RuntimeError("El modelo no fue cargado.")

        for start in range(0, len(crops), self.batch_size):
            batch = crops[start : start + self.batch_size]
            images = [crop.image for crop in batch]
            results = self.predict_batch(images)
            for crop, result in zip(batch, results):
                boxes = getattr(result, "boxes", [])
                for box in boxes:
                    class_id = int(safe_float(box.cls[0]))
                    confidence = safe_float(box.conf[0])
                    if confidence < self.confidence or class_id < 0 or class_id >= len(self.prompts):
                        continue

                    prompt = self.prompts[class_id]
                    target = self.prompt_to_target.get(normalize_key(prompt))
                    if target is None:
                        continue

                    bx1, by1, bx2, by2 = safe_xyxy(box.xyxy[0])
                    x1 = bx1 / crop.scale + crop.x0
                    y1 = by1 / crop.scale + crop.y0
                    x2 = bx2 / crop.scale + crop.x0
                    y2 = by2 / crop.scale + crop.y0
                    ix1, iy1, ix2, iy2 = clamp_box(x1, y1, x2, y2, image_width, image_height)
                    box_width = ix2 - ix1
                    box_height = iy2 - iy1
                    area_ratio = (box_width * box_height) / max(1, image_width * image_height)
                    detections.append(
                        Detection(
                            image=image_name,
                            target=target.label,
                            prompt=prompt,
                            group=target.group,
                            priority=target.priority,
                            confidence=confidence,
                            x1=ix1,
                            y1=iy1,
                            x2=ix2,
                            y2=iy2,
                            width_px=box_width,
                            height_px=box_height,
                            area_ratio=area_ratio,
                            source_pass=crop.pass_name,
                            **metadata,
                        )
                    )
        return detections

    def predict_batch(self, images: list[np.ndarray]) -> list[Any]:
        if self.model is None:
            raise RuntimeError("El modelo no fue cargado.")

        with contextlib.suppress(Exception):
            import torch

            with torch.no_grad():
                return self.model.predict(
                    source=images,
                    conf=self.confidence,
                    imgsz=self.imgsz,
                    batch=len(images),
                    verbose=False,
                )

        try:
            return self.model.predict(
                source=images,
                conf=self.confidence,
                imgsz=self.imgsz,
                batch=len(images),
                verbose=False,
            )
        except Exception:
            results = []
            for image in images:
                results.extend(self.model.predict(source=image, conf=self.confidence, imgsz=self.imgsz, verbose=False))
            return results

    def apply_nms(self, detections: list[Detection]) -> list[Detection]:
        if not detections:
            return []

        by_prompt: dict[str, list[Detection]] = defaultdict(list)
        for detection in detections:
            by_prompt[detection.prompt].append(detection)

        final: list[Detection] = []
        for prompt, prompt_detections in by_prompt.items():
            boxes = [
                [det.x1, det.y1, max(1, det.x2 - det.x1), max(1, det.y2 - det.y1)]
                for det in prompt_detections
            ]
            scores = [det.confidence for det in prompt_detections]
            indices = cv2.dnn.NMSBoxes(boxes, scores, score_threshold=0.0, nms_threshold=self.nms_iou)
            if indices is None or len(indices) == 0:
                continue
            flat_indices = np.array(indices).reshape(-1)
            final.extend(prompt_detections[int(i)] for i in flat_indices)

        final.sort(key=lambda det: (det.image, det.target, -det.confidence))
        return final

    def draw_detections(self, image: np.ndarray, image_name: str, detections: list[Detection]) -> np.ndarray:
        height, width = image.shape[:2]
        thickness = max(2, int(round(max(width, height) / 900)))
        font_scale = max(0.55, min(1.25, max(width, height) / 1800))
        target_by_label = {target.label: target for target in self.targets}

        overlay = image.copy()
        cv2.rectangle(overlay, (0, 0), (width, max(46, int(58 * font_scale))), (20, 34, 50), -1)
        cv2.addWeighted(overlay, 0.88, image, 0.12, 0, image)
        title = f"{image_name} | Hallazgos: {len(detections)}"
        cv2.putText(
            image,
            title[:90],
            (16, max(30, int(36 * font_scale))),
            cv2.FONT_HERSHEY_SIMPLEX,
            font_scale,
            (255, 255, 255),
            thickness,
            cv2.LINE_AA,
        )

        for detection in detections:
            target = target_by_label.get(detection.target)
            color = hex_to_bgr(target.color if target else "#2563eb")
            cv2.rectangle(image, (detection.x1, detection.y1), (detection.x2, detection.y2), color, thickness)
            label = f"{detection.target} {detection.confidence:.2f}"
            (label_width, label_height), baseline = cv2.getTextSize(
                label,
                cv2.FONT_HERSHEY_SIMPLEX,
                font_scale,
                thickness,
            )
            x1 = detection.x1
            y1 = max(label_height + baseline + 6, detection.y1)
            cv2.rectangle(
                image,
                (x1, y1 - label_height - baseline - 8),
                (min(width - 1, x1 + label_width + 10), y1 + 4),
                color,
                -1,
            )
            cv2.putText(
                image,
                label,
                (x1 + 5, y1 - baseline - 2),
                cv2.FONT_HERSHEY_SIMPLEX,
                font_scale,
                (255, 255, 255),
                max(1, thickness - 1),
                cv2.LINE_AA,
            )
        return image

    def detections_dataframe(self) -> pd.DataFrame:
        if not self.detections:
            return pd.DataFrame(
                columns=[
                    "image",
                    "target",
                    "group",
                    "priority",
                    "confidence",
                    "x1",
                    "y1",
                    "x2",
                    "y2",
                    "source_pass",
                ]
            )
        return pd.DataFrame(asdict(det) for det in self.detections)

    def images_dataframe(self) -> pd.DataFrame:
        if not self.image_summaries:
            return pd.DataFrame()
        return pd.DataFrame(asdict(summary) for summary in self.image_summaries)

    def encounters_dataframe(self) -> pd.DataFrame:
        df = self.detections_dataframe()
        if df.empty:
            return pd.DataFrame(
                columns=[
                    "hallazgo",
                    "cantidad",
                    "imagenes_con_hallazgo",
                    "confianza_promedio",
                    "confianza_maxima",
                    "grupo",
                    "prioridad",
                    "accion_sugerida",
                ]
            )

        target_action = {target.label: target.action for target in self.targets}
        rows = []
        for target, group in df.groupby("target"):
            rows.append(
                {
                    "hallazgo": target,
                    "cantidad": int(len(group)),
                    "imagenes_con_hallazgo": int(group["image"].nunique()),
                    "confianza_promedio": round(float(group["confidence"].mean()), 3),
                    "confianza_maxima": round(float(group["confidence"].max()), 3),
                    "grupo": str(group["group"].iloc[0]),
                    "prioridad": str(group["priority"].iloc[0]),
                    "accion_sugerida": target_action.get(target, "Validar visualmente y registrar decision."),
                }
            )
        result = pd.DataFrame(rows)
        priority_order = {"Alta": 0, "Media": 1, "Baja": 2}
        result["_priority_order"] = result["prioridad"].map(priority_order).fillna(9)
        result = result.sort_values(["_priority_order", "cantidad", "hallazgo"], ascending=[True, False, True])
        return result.drop(columns=["_priority_order"])

    def config_dataframe(self) -> pd.DataFrame:
        rows = [
            ("fecha_inicio", self.started_at.strftime("%Y-%m-%d %H:%M:%S")),
            ("fecha_fin", self.finished_at.strftime("%Y-%m-%d %H:%M:%S") if self.finished_at else ""),
            ("input_folder", str(self.input_folder)),
            ("output_folder", str(self.output_folder)),
            ("modelo", self.model_name),
            ("confianza", self.confidence),
            ("perfil_busqueda", self.search_profile),
            ("pasadas", ", ".join(search_pass.name for search_pass in self.search_passes)),
            ("nms_iou", self.nms_iou),
            ("max_image_side", self.max_image_side),
            ("batch_size", self.batch_size),
            ("imgsz", self.imgsz),
            ("objetivos", ", ".join(f"{target.label} -> {target.prompt}" for target in self.targets)),
        ]
        return pd.DataFrame(rows, columns=["parametro", "valor"])

    def export_all(self) -> dict[str, Path]:
        artifacts = {
            "detecciones_csv": self.write_csv(self.detections_dataframe(), "detecciones_detalle.csv"),
            "imagenes_csv": self.write_csv(self.images_dataframe(), "imagenes_resumen.csv"),
            "encuentros_csv": self.write_csv(self.encounters_dataframe(), "encuentros_resumen.csv"),
            "excel": self.write_excel(),
            "pdf": self.write_pdf(),
        }
        artifacts["zip"] = self.write_zip(artifacts)
        return artifacts

    def write_csv(self, dataframe: pd.DataFrame, filename: str) -> Path:
        path = self.tables_folder / filename
        dataframe.to_csv(path, index=False, encoding="utf-8-sig")
        return path

    def write_excel(self) -> Path:
        path = self.output_folder / "reporte_inspeccion_visual.xlsx"
        with pd.ExcelWriter(path, engine="openpyxl") as writer:
            self.encounters_dataframe().to_excel(writer, index=False, sheet_name="Encuentros")
            self.detections_dataframe().to_excel(writer, index=False, sheet_name="Detalle_detecciones")
            self.images_dataframe().to_excel(writer, index=False, sheet_name="Fotos")
            self.config_dataframe().to_excel(writer, index=False, sheet_name="Configuracion")
            self.recommendations_dataframe().to_excel(writer, index=False, sheet_name="Recomendaciones")

        self.style_excel(path)
        return path

    def style_excel(self, path: Path) -> None:
        try:
            from openpyxl import load_workbook
            from openpyxl.styles import Alignment, Font, PatternFill
        except Exception:
            return

        workbook = load_workbook(path)
        header_fill = PatternFill("solid", fgColor="1F4E5F")
        header_font = Font(color="FFFFFF", bold=True)
        soft_fill = PatternFill("solid", fgColor="EAF4F4")

        for sheet in workbook.worksheets:
            sheet.freeze_panes = "A2"
            for cell in sheet[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center")
            for row in sheet.iter_rows(min_row=2):
                for cell in row:
                    cell.alignment = Alignment(vertical="top", wrap_text=True)
                if row[0].row % 2 == 0:
                    for cell in row:
                        cell.fill = soft_fill
            for column_cells in sheet.columns:
                column = column_cells[0].column_letter
                max_length = max(len(str(cell.value or "")) for cell in column_cells[:80])
                sheet.column_dimensions[column].width = min(max(12, max_length + 2), 48)

        workbook.save(path)

    def recommendations_dataframe(self) -> pd.DataFrame:
        encounters = self.encounters_dataframe()
        if encounters.empty:
            return pd.DataFrame(
                [
                    {
                        "tema": "Resultado",
                        "grupo": "Gestion",
                        "prioridad": "Media",
                        "cantidad": "",
                        "recomendacion": "No se detectaron los objetivos configurados. Revisar muestra y sensibilidad si era esperable encontrarlos.",
                    }
                ]
            )

        rows = []
        for _, row in encounters.iterrows():
            rows.append(
                {
                    "tema": row["hallazgo"],
                    "grupo": row["grupo"],
                    "prioridad": row["prioridad"],
                    "cantidad": row["cantidad"],
                    "recomendacion": row["accion_sugerida"],
                }
            )
        rows.append(
            {
                "tema": "Uso del reporte",
                "grupo": "Gestion",
                "prioridad": "Media",
                "cantidad": "",
                "recomendacion": "La IA acelera el relevamiento, pero cada hallazgo debe validarse por responsable tecnico o SyH antes de emitir conformidad.",
            }
        )
        return pd.DataFrame(rows)

    def write_pdf(self) -> Path:
        try:
            from reportlab.lib import colors
            from reportlab.lib.enums import TA_CENTER
            from reportlab.lib.pagesizes import A4, landscape
            from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
            from reportlab.lib.units import mm
            from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
        except ImportError as exc:
            raise RuntimeError("Falta reportlab. Instala con: pip install reportlab") from exc

        pdf_path = self.output_folder / "reporte_inspeccion_visual.pdf"
        doc = SimpleDocTemplate(
            str(pdf_path),
            pagesize=landscape(A4),
            leftMargin=12 * mm,
            rightMargin=12 * mm,
            topMargin=12 * mm,
            bottomMargin=10 * mm,
        )

        styles = getSampleStyleSheet()
        styles.add(
            ParagraphStyle(
                "ReportTitle",
                parent=styles["Title"],
                fontName="Helvetica-Bold",
                fontSize=24,
                textColor=colors.HexColor("#16323F"),
                alignment=TA_CENTER,
                spaceAfter=10,
            )
        )
        styles.add(
            ParagraphStyle(
                "SectionTitle",
                parent=styles["Heading2"],
                fontSize=14,
                textColor=colors.HexColor("#1F4E5F"),
                spaceBefore=8,
                spaceAfter=8,
            )
        )
        styles.add(
            ParagraphStyle(
                "Small",
                parent=styles["Normal"],
                fontSize=8,
                leading=10,
                textColor=colors.HexColor("#334155"),
            )
        )

        story: list[Any] = []
        story.append(Paragraph("Reporte de Inspeccion Visual", styles["ReportTitle"]))
        story.append(
            Paragraph(
                "Oficina tecnica + Seguridad e Higiene | Deteccion multiobjeto con IA | "
                f"{datetime.now().strftime('%d/%m/%Y %H:%M')}",
                styles["Normal"],
            )
        )
        story.append(Spacer(1, 10))

        total_images = len(self.image_summaries)
        images_with_findings = sum(1 for summary in self.image_summaries if summary.total_detections > 0)
        total_detections = len(self.detections)
        summary_data = [
            ["Fotos procesadas", "Fotos con hallazgos", "Detecciones", "Objetivos buscados", "Perfil"],
            [
                str(total_images),
                str(images_with_findings),
                str(total_detections),
                str(len(self.targets)),
                self.search_profile,
            ],
        ]
        summary_table = Table(summary_data, colWidths=[38 * mm, 42 * mm, 35 * mm, 42 * mm, 42 * mm])
        summary_table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4E5F")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                    ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                    ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#CBD5E1")),
                    ("BACKGROUND", (0, 1), (-1, 1), colors.HexColor("#EAF4F4")),
                    ("FONTSIZE", (0, 0), (-1, -1), 10),
                ]
            )
        )
        story.append(summary_table)

        story.append(Paragraph("Encuentros y cantidades", styles["SectionTitle"]))
        encounters = self.encounters_dataframe()
        if encounters.empty:
            story.append(Paragraph("No se detectaron los objetivos configurados.", styles["Normal"]))
        else:
            encounter_rows = [["Hallazgo", "Cantidad", "Fotos", "Conf. prom.", "Grupo", "Prioridad"]]
            for _, row in encounters.iterrows():
                encounter_rows.append(
                    [
                        html.escape(str(row["hallazgo"])),
                        str(row["cantidad"]),
                        str(row["imagenes_con_hallazgo"]),
                        str(row["confianza_promedio"]),
                        html.escape(str(row["grupo"])),
                        html.escape(str(row["prioridad"])),
                    ]
                )
            table = Table(encounter_rows, colWidths=[52 * mm, 24 * mm, 22 * mm, 28 * mm, 42 * mm, 30 * mm])
            table.setStyle(self.default_table_style(colors))
            story.append(table)

        story.append(Paragraph("Lectura tecnica sugerida", styles["SectionTitle"]))
        recs = self.recommendations_dataframe().head(8)
        rec_rows = [["Tema", "Prioridad", "Recomendacion"]]
        for _, row in recs.iterrows():
            rec_rows.append(
                [
                    html.escape(str(row.get("tema", ""))),
                    html.escape(str(row.get("prioridad", ""))),
                    Paragraph(html.escape(str(row.get("recomendacion", ""))), styles["Small"]),
                ]
            )
        rec_table = Table(rec_rows, colWidths=[48 * mm, 27 * mm, 122 * mm])
        rec_table.setStyle(self.default_table_style(colors))
        story.append(rec_table)

        story.append(PageBreak())
        story.append(Paragraph("Detalle por foto", styles["SectionTitle"]))
        photo_rows = [["Foto", "Estado", "Hallazgos", "GPS/Altitud"]]
        for summary in self.image_summaries:
            gps = ""
            if summary.gps_lat is not None and summary.gps_lon is not None:
                gps = f"{summary.gps_lat:.6f}, {summary.gps_lon:.6f}"
            if summary.gps_alt_m is not None:
                gps = (gps + " | " if gps else "") + f"{summary.gps_alt_m:.1f} m"
            photo_rows.append(
                [
                    html.escape(summary.image[:48]),
                    summary.status,
                    Paragraph(html.escape(summary.findings), styles["Small"]),
                    gps or "N/D",
                ]
            )
        photo_table = Table(photo_rows, colWidths=[58 * mm, 34 * mm, 88 * mm, 42 * mm], repeatRows=1)
        photo_table.setStyle(self.default_table_style(colors))
        story.append(photo_table)

        annotated = [summary for summary in self.image_summaries if summary.annotated_path]
        if annotated:
            story.append(PageBreak())
            story.append(Paragraph("Anexo fotografico", styles["SectionTitle"]))
            for summary in annotated:
                story.append(Paragraph(html.escape(summary.image), styles["Heading3"]))
                story.append(Paragraph(html.escape(summary.findings), styles["Small"]))
                image_flowable = self.reportlab_image(summary.annotated_path, max_width=240 * mm, max_height=128 * mm)
                if image_flowable:
                    story.append(image_flowable)
                    story.append(Spacer(1, 8))

        story.append(Spacer(1, 6))
        story.append(
            Paragraph(
                "Nota: este reporte es una ayuda de deteccion y trazabilidad. No reemplaza la inspeccion de un "
                "profesional habilitado ni el criterio del responsable de Higiene y Seguridad.",
                styles["Small"],
            )
        )

        def page_footer(canvas: Any, doc_obj: Any) -> None:
            canvas.saveState()
            canvas.setFont("Helvetica", 7)
            canvas.setFillColor(colors.HexColor("#64748B"))
            canvas.drawRightString(285 * mm, 7 * mm, f"Pagina {doc_obj.page}")
            canvas.drawString(12 * mm, 7 * mm, "Reporte generado por Buscador Visual Oficina Tecnica")
            canvas.restoreState()

        doc.build(story, onFirstPage=page_footer, onLaterPages=page_footer)
        return pdf_path

    @staticmethod
    def default_table_style(colors_module: Any) -> Any:
        from reportlab.platypus import TableStyle

        return TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors_module.HexColor("#1F4E5F")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors_module.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("ALIGN", (0, 0), (-1, 0), "CENTER"),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("GRID", (0, 0), (-1, -1), 0.35, colors_module.HexColor("#CBD5E1")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors_module.white, colors_module.HexColor("#F8FAFC")]),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                ("LEADING", (0, 0), (-1, -1), 9),
            ]
        )

    @staticmethod
    def reportlab_image(path: str | Path, max_width: float, max_height: float) -> Any | None:
        try:
            from reportlab.platypus import Image as RLImage

            with Image.open(path) as image:
                width, height = image.size
            scale = min(max_width / width, max_height / height)
            return RLImage(str(path), width=width * scale, height=height * scale)
        except Exception:
            return None

    def write_zip(self, artifacts: dict[str, Path]) -> Path:
        zip_path = self.output_folder / "paquete_reporte_visual.zip"
        with zipfile.ZipFile(zip_path, "w", compression=zipfile.ZIP_DEFLATED) as zf:
            for name, path in artifacts.items():
                if name == "zip" or not Path(path).exists():
                    continue
                zf.write(path, arcname=Path(path).relative_to(self.output_folder))
            for image_path in sorted(self.annotated_folder.glob("*.jpg")):
                zf.write(image_path, arcname=image_path.relative_to(self.output_folder))
        return zip_path


def run_inspection(
    input_folder: str | Path,
    output_folder: str | Path = "output_reporte_visual",
    preset_labels: Iterable[str] | None = DEFAULT_PRESET_LABELS,
    extra_terms: str | Iterable[str] | None = None,
    confidence: float = 0.08,
    search_profile: str = "equilibrada",
    batch_size: int = 8,
    model_name: str = "yolov8s-world.pt",
) -> tuple[VisualInspectionTool, dict[str, Path]]:
    targets = build_targets(preset_labels=preset_labels, extra_terms=extra_terms)
    tool = VisualInspectionTool(
        input_folder=input_folder,
        output_folder=output_folder,
        targets=targets,
        confidence=confidence,
        search_profile=search_profile,
        batch_size=batch_size,
        model_name=model_name,
    )
    tool.process()
    artifacts = tool.export_all()
    return tool, artifacts


def copy_uploaded_files(uploaded: dict[str, bytes], input_folder: str | Path) -> None:
    folder = Path(input_folder)
    folder.mkdir(parents=True, exist_ok=True)
    for filename in uploaded.keys():
        source = Path(filename)
        if source.exists() and source.suffix.lower() in IMAGE_EXTENSIONS:
            shutil.move(str(source), str(folder / source.name))


def _running_in_notebook() -> bool:
    return "ipykernel" in sys.modules or "google.colab" in sys.modules


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Buscador visual multiobjeto para reportes de obra.")
    parser.add_argument("--input", required=True, help="Carpeta con fotos.")
    parser.add_argument("--output", default="output_reporte_visual", help="Carpeta de salida.")
    parser.add_argument("--buscar", default="", help="Objetos extra separados por coma: cono, cartel, casco, autos.")
    parser.add_argument(
        "--preset",
        action="append",
        default=[],
        help="Etiqueta de catalogo a incluir. Puede repetirse. Si se omite, usa el preset recomendado.",
    )
    parser.add_argument("--confianza", type=float, default=0.08, help="Confianza minima de deteccion.")
    parser.add_argument(
        "--perfil",
        choices=sorted(SEARCH_PROFILES),
        default="equilibrada",
        help="rapida, equilibrada o profunda.",
    )
    parser.add_argument("--batch-size", type=int, default=8, help="Cantidad de sectores por lote.")
    parser.add_argument("--modelo", default="yolov8s-world.pt", help="Modelo YOLO-World.")
    args = parser.parse_args(argv)

    presets = args.preset or list(DEFAULT_PRESET_LABELS)
    _, artifacts = run_inspection(
        input_folder=args.input,
        output_folder=args.output,
        preset_labels=presets,
        extra_terms=args.buscar,
        confidence=args.confianza,
        search_profile=args.perfil,
        batch_size=args.batch_size,
        model_name=args.modelo,
    )
    print(json.dumps({key: str(path) for key, path in artifacts.items()}, indent=2, ensure_ascii=False))
    return 0


if __name__ == "__main__" and not _running_in_notebook():
    raise SystemExit(main())
